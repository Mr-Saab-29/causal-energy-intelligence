from __future__ import annotations

import io

import pandas as pd
import pytest
from openpyxl import Workbook

from src.data.pre_treatment_covariates import (
    assemble_hourly_covariates,
    build_pre_treatment_readiness,
    load_pre_treatment_contract,
    parse_eex_auction_workbook,
    parse_world_bank_fuel_workbook,
)


def workbook_bytes(workbook: Workbook) -> bytes:
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_world_bank_parser_reads_european_gas_and_coal_series() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Monthly Prices"
    sheet.append([])
    sheet.append([])
    sheet.append([])
    sheet.append([])
    sheet.append(["Month", "Natural gas, Europe", "Coal, Australian"])
    sheet.append(["", "$/mmbtu", "$/mt"])
    sheet.append(["2023M01", 20.5, 310.0])

    result = parse_world_bank_fuel_workbook(
        workbook_bytes(workbook), load_pre_treatment_contract()
    )

    assert result.to_dict(orient="records") == [
        {
            "gas_price_usd_mmbtu": 20.5,
            "coal_price_usd_mt": 310.0,
            "reference_month": pd.Period("2023-01", freq="M"),
        }
    ]


def test_eex_parser_keeps_only_completed_eu_primary_auctions() -> None:
    workbook = Workbook()
    sheet = workbook.active
    for _ in range(5):
        sheet.append([])
    sheet.append(["Time", "Contract", "Status", "Auction Price €/tCO2", "Zone"])
    sheet.append([pd.Timestamp("2023-01-09 11:00"), "T3PA", "successful", 82.5, "EU"])
    sheet.append([pd.Timestamp("2023-01-10 11:00"), "T3PA", "cancelled", 90.0, "EU"])
    sheet.append([pd.Timestamp("2023-01-11 11:00"), "T3PA", "successful", 70.0, "DE"])

    result = parse_eex_auction_workbook(
        workbook_bytes(workbook), load_pre_treatment_contract()
    )

    assert len(result) == 1
    assert result.iloc[0]["available_at_utc"] == pd.Timestamp(
        "2023-01-09 10:00", tz="UTC"
    )
    assert result.iloc[0]["eua_auction_price_eur_tco2"] == pytest.approx(82.5)


def test_eex_parser_supports_archived_auction_name_zone() -> None:
    workbook = Workbook()
    sheet = workbook.active
    for _ in range(5):
        sheet.append([])
    sheet.append(
        ["Time", "Auction Name", "Contract", "Status", "Auction Price €/tCO2"]
    )
    sheet.append(
        [
            pd.Timestamp("2023-01-09 11:00"),
            "Auction 4. Period CAP3 EU",
            "T3PA",
            "successful",
            82.5,
        ]
    )
    sheet.append(
        [
            pd.Timestamp("2023-01-10 11:00"),
            "Auction 4. Period DE",
            "T3PA",
            "successful",
            80.0,
        ]
    )

    result = parse_eex_auction_workbook(
        workbook_bytes(workbook), load_pre_treatment_contract()
    )

    assert len(result) == 1
    assert result.iloc[0]["eua_auction_price_eur_tco2"] == pytest.approx(82.5)


def test_hourly_alignment_uses_only_values_available_before_decision_time() -> None:
    contract = load_pre_treatment_contract()
    start = pd.Timestamp("2023-03-01 00:00", tz="UTC")
    end = pd.Timestamp("2023-03-01 03:00", tz="UTC")
    weather = pd.DataFrame(
        {
            "timestamp_utc": pd.date_range(start, end, freq="h", inclusive="left"),
            "weather_temperature_forecast_c_24h": [8.0, 9.0, 10.0],
            "weather_region_count": [12, 12, 12],
        }
    )
    fuel = pd.DataFrame(
        {
            "reference_month": [pd.Period("2023-01", freq="M")],
            "gas_price_usd_mmbtu": [20.0],
            "coal_price_usd_mt": [300.0],
        }
    )
    carbon = pd.DataFrame(
        {
            "available_at_utc": [start, start + pd.Timedelta(hours=1)],
            "eua_auction_price_eur_tco2": [80.0, 81.0],
        }
    )
    storage = pd.DataFrame(
        {
            "observed_at_utc": [start - pd.Timedelta(days=7)],
            "hydro_storage_mwh": [2_000_000.0],
        }
    )

    result = assemble_hourly_covariates(
        start, end, contract, weather, fuel, carbon, storage
    )

    assert pd.isna(result.loc[0, "eua_auction_price_eur_tco2"])
    assert result.loc[1, "eua_auction_price_eur_tco2"] == pytest.approx(80.0)
    assert pd.isna(result.loc[0, "hydro_storage_mwh_lag_1w"])
    assert result.loc[1, "hydro_storage_mwh_lag_1w"] == pytest.approx(2_000_000.0)
    assert result["gas_price_usd_mmbtu_lag_2m"].eq(20.0).all()


def test_readiness_passes_complete_temporally_safe_window() -> None:
    contract = load_pre_treatment_contract()
    start = pd.Timestamp("2023-03-01 00:00", tz="UTC")
    hours = pd.date_range(start, periods=48, freq="h")
    frame = pd.DataFrame(
        {
            "timestamp_utc": hours,
            "weather_temperature_forecast_c_24h": 8.0,
            "gas_price_usd_mmbtu_lag_2m": 20.0,
            "coal_price_usd_mt_lag_2m": 300.0,
            "eua_auction_price_eur_tco2": 80.0,
            "hydro_storage_mwh_lag_1w": 2_000_000.0,
            "fuel_reference_month": pd.Timestamp("2023-01-01"),
            "carbon_auction_at_utc": start - pd.Timedelta(hours=1),
            "hydro_storage_observed_at_utc": start - pd.Timedelta(days=8),
        }
    )

    report = build_pre_treatment_readiness(frame, contract)

    assert report["status"] == "pass"
    assert report["pre_treatment_covariates_ready"] is True
    assert not report["critical_issues"]
