"""Build the interview methodology handbook from its Markdown source."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE_PATH = ROOT / "docs/causal_energy_methodology_interview_handbook.md"
OUTPUT_PATH = ROOT / "docs/causal_energy_methodology_interview_handbook.docx"

FONT = "Liberation Sans"
MONO_FONT = "Liberation Mono"
INK = "17211D"
MUTED = "5E6B65"
TABLE_HEADER = "315A4B"
TABLE_ALT = "EDF4F1"
TABLE_BORDER = "D9D9D9"
CODE_FILL = "F3F5F4"


def set_run_font(run, name: str = FONT) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.7)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12

    style_specs = {
        "Title": (28, 0, 12),
        "Subtitle": (14, 0, 16),
        "Heading 1": (17, 14, 7),
        "Heading 2": (13, 11, 5),
        "Heading 3": (11.5, 9, 4),
    }
    for style_name, (size, before, after) in style_specs.items():
        style = styles[style_name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    document.core_properties.title = "Causal Energy Intelligence Methodology Handbook"
    document.core_properties.subject = "Forecasting recommendations causal inference and operations"
    document.core_properties.author = "Causal Energy Intelligence Project"


def add_footer(document: Document) -> None:
    for section in document.sections:
        footer = section.footer
        table = footer.add_table(rows=1, cols=2, width=Inches(6.86))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.columns[0].width = Inches(5.8)
        table.columns[1].width = Inches(1.06)
        remove_table_borders(table)

        left = table.cell(0, 0).paragraphs[0]
        left.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = left.add_run("Causal Energy Intelligence Methodology Handbook")
        set_run_font(run)
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)

        right = table.cell(0, 1).paragraphs[0]
        right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = right.add_run("Page ")
        set_run_font(run)
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        right._p.append(field)


def remove_table_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)


def add_page_break(document: Document) -> None:
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_inline_runs(paragraph, text: str) -> None:
    pattern = re.compile(r"(\*\*.+?\*\*|`.+?`)")
    for part in pattern.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
            set_run_font(run)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, MONO_FONT)
            run.font.size = Pt(9.5)
        else:
            run = paragraph.add_run(part)
            set_run_font(run)


def add_body_paragraph(document: Document, text: str, style: str | None = None):
    paragraph = document.add_paragraph(style=style)
    add_inline_runs(paragraph, text)
    paragraph.paragraph_format.widow_control = True
    return paragraph


def add_numbered_paragraph(document: Document, number: int, text: str):
    paragraph = add_body_paragraph(document, f"{number}. {text}")
    paragraph.paragraph_format.left_indent = Inches(0.24)
    paragraph.paragraph_format.first_line_indent = Inches(-0.24)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    for run in paragraph.runs:
        run.font.size = Pt(10)
    return paragraph


def add_heading(document: Document, text: str, level: int):
    paragraph = document.add_heading(text, level=level)
    for run in paragraph.runs:
        set_run_font(run)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_code_block(document: Document, lines: list[str]) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.28)
    paragraph.paragraph_format.right_indent = Inches(0.28)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), CODE_FILL)
    paragraph._p.get_or_add_pPr().append(shading)
    run = paragraph.add_run("\n".join(lines))
    set_run_font(run, MONO_FONT)
    run.font.size = Pt(8.8)


def set_cell_shading(cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_margins(cell, top: int = 90, start: int = 110, bottom: int = 90, end: int = 110):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = margins.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), TABLE_BORDER)
        borders.append(element)


def repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def prevent_row_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    properties.append(OxmlElement("w:cantSplit"))


def add_markdown_table(document: Document, rows: list[list[str]]) -> None:
    if len(rows) < 2:
        return
    header = rows[0]
    body = rows[2:] if all(set(value) <= {"-", ":"} for value in rows[1]) else rows[1:]
    table = document.add_table(rows=1, cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])

    for index, value in enumerate(header):
        cell = table.rows[0].cells[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, TABLE_HEADER)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(value.strip())
        set_run_font(run)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)

    for row_index, values in enumerate(body):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for index, value in enumerate(values):
            cell = cells[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index % 2 == 1:
                set_cell_shading(cell, TABLE_ALT)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
                if index > 0 and len(value.strip()) < 24
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            add_inline_runs(paragraph, value.strip())
            for run in paragraph.runs:
                run.font.size = Pt(8.8)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def parse_table_line(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def collect_contents(lines: list[str]) -> list[str]:
    headings = []
    for line in lines:
        if line.startswith("## "):
            text = line[3:].strip()
            if text != "Forecasting Recommendations Causal Inference and Production Operations":
                headings.append(text)
    return headings


def build_document(source_path: Path = SOURCE_PATH, output_path: Path = OUTPUT_PATH) -> Path:
    lines = source_path.read_text(encoding="utf-8").splitlines()
    document = Document()
    configure_document(document)
    add_footer(document)
    contents = collect_contents(lines)

    index = 0
    in_code = False
    code_lines: list[str] = []
    cover = True
    contents_added = False
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped == "```text":
            in_code = True
            code_lines = []
            index += 1
            continue
        if in_code:
            if stripped == "```":
                add_code_block(document, code_lines)
                in_code = False
            else:
                code_lines.append(line)
            index += 1
            continue
        if stripped == "<!-- pagebreak -->":
            add_page_break(document)
            cover = False
            if not contents_added:
                add_heading(document, "Contents", 1)
                for number, heading in enumerate(contents, start=1):
                    add_numbered_paragraph(document, number, heading)
                add_page_break(document)
                contents_added = True
            index += 1
            continue
        if line.startswith("| "):
            table_rows: list[list[str]] = []
            while index < len(lines) and lines[index].startswith("|"):
                table_rows.append(parse_table_line(lines[index]))
                index += 1
            add_markdown_table(document, table_rows)
            continue
        if line.startswith("# "):
            paragraph = document.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Inches(1.0)
            add_inline_runs(paragraph, line[2:].strip())
        elif line.startswith("## "):
            text = line[3:].strip()
            if cover:
                paragraph = document.add_paragraph(style="Subtitle")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline_runs(paragraph, text)
            else:
                heading = add_heading(document, text, 1)
                if text == "Research References":
                    heading.paragraph_format.page_break_before = True
        elif line.startswith("### "):
            add_heading(document, line[4:].strip(), 2)
        elif line.startswith("#### "):
            add_heading(document, line[5:].strip(), 3)
        elif match := re.match(r"^(\d+)\. (.+)", stripped):
            add_numbered_paragraph(document, int(match.group(1)), match.group(2))
        elif stripped.startswith("- "):
            add_body_paragraph(document, stripped[2:], "List Bullet")
        elif stripped:
            paragraph = add_body_paragraph(document, stripped)
            if cover and stripped.startswith("Version "):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_after = Pt(16)
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor.from_string(MUTED)
            elif cover:
                paragraph.paragraph_format.left_indent = Inches(0.35)
                paragraph.paragraph_format.right_indent = Inches(0.35)
        index += 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path


if __name__ == "__main__":
    path = build_document()
    print(path)
